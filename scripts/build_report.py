"""Legge i Riepiloghi mensili C1 (SIAE, cassa 18Tickets) e aggiorna:
- data/HorrorDays_Vendite.xlsx (report Excel, con storico conservato)
- site/data.json (dati per il sito del report)
Uso: python scripts/build_report.py [--c1 data/c1] [--xlsx data/HorrorDays_Vendite.xlsx] [--json site/data.json]
"""
import re, glob, sys, datetime as dt, os, json, subprocess, argparse
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList

ap = argparse.ArgumentParser()
ap.add_argument("--c1", default="data/c1")
ap.add_argument("--xlsx", default="data/HorrorDays_Vendite.xlsx")
ap.add_argument("--json", default="site/data.json")
ap.add_argument("--xlsx-copy", default="site/HorrorDays_Vendite.xlsx", help="copia scaricabile dal sito ('' per non pubblicarla)")
args = ap.parse_args()
OUT = PREV = args.xlsx
SNAP = dt.datetime.now(ZoneInfo("Europe/Rome")).replace(tzinfo=None, second=0, microsecond=0)

# Ogni titolo d'accesso è un'auto (la capienza del C1 è in auto). Le persone dipendono dal prezzo:
# 24 € = 2 persone, 36 € = 3, 48 € = 4, cioè 12 € a persona. Prezzi diversi si definiscono
# nel foglio Tariffe (colonna "Persone per auto"), altrimenti finiscono tra le anomalie.
PERSONE_PER_PREZZO = {24.0: 2, 36.0: 3, 48.0: 4}
EURO_A_PERSONA = 12.0
TIPI = sorted(set(PERSONE_PER_PREZZO.values()))  # 2, 3, 4

pdfs = sorted(glob.glob(os.path.join(args.c1, "*.pdf")) + glob.glob(os.path.join(args.c1, "*.PDF")))
if not pdfs:
    sys.exit("Nessun PDF C1 trovato in " + args.c1)
num = r"(-?[\d]+\.\d{2})"
row_re = re.compile(r"^\s*(\S+.*?)\s{2,}(\d+)\s+(\S+)\s+" + num + r"\s+(\d+)\s+" + num + r"\s+" + num + r"\s+" + num + r"\s+" + num + r"\s+" + num + r"\s+" + num + r"\s+(\d+)\s*$")
tot_re = re.compile(r"TOTALE GENERALE\s+(\d+)\s+" + num)

shows, sales, total_pages, pdf_tot_gen = [], [], 0, 0.0
for p in pdfs:
    info = subprocess.run(["pdfinfo", p], capture_output=True, text=True).stdout
    npg = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    first = subprocess.run(["pdftotext", "-layout", "-f", "1", "-l", "1", p, "-"], capture_output=True, text=True).stdout
    mm = re.search(r"Riepilogo mensile del\s+(\d{2}/\d{4})", first)
    mese = mm.group(1) if mm else os.path.basename(p)
    if True:
        for pg in range(1, npg + 1):
            total_pages += 1
            t = subprocess.run(["pdftotext", "-layout", "-f", str(pg), "-l", str(pg), p, "-"], capture_output=True, text=True).stdout
            d = re.search(r"DATA EVENTO\s+(\d{2}/\d{2}/\d{4})", t).group(1)
            o = re.search(r"ORA INIZIO\s+(\d{2}:\d{2})", t).group(1)
            tit = re.search(r"MANIFESTAZIONE\s+(.+?)\s{2,}AUTORE", t).group(1).strip()
            loc = re.search(r"DENOMINAZIONE LOCALE:\s*(.+?)\s{2,}", t).group(1).strip()
            cod = re.search(r"CODICE LOCALE:\s*(\d+)", t).group(1)
            tg = tot_re.search(t)
            pdf_tot_gen += float(tg.group(2))
            date = dt.datetime.strptime(d, "%d/%m/%Y").date()
            hh, mm = map(int, o.split(":"))
            sid = date.strftime("%Y%m%d") + "_" + o.replace(":", "")
            cap = None
            for line in t.splitlines():
                if "TOTALE" in line:
                    continue
                m = row_re.match(line)
                if m:
                    g = m.groups()
                    cap = int(g[1])
                    sales.append(dict(id=sid, date=date, time=dt.time(hh, mm), settore=g[0].strip(), cap=cap,
                                      code=g[2], prezzo=float(g[3]), n=int(g[4]), lordo=float(g[5]),
                                      prev=float(g[6]), impint=float(g[7]), intr=float(g[8]),
                                      impiva=float(g[9]), iva=float(g[10]), ann=int(g[11]), mese=mese))
            shows.append(dict(id=sid, date=date, time=dt.time(hh, mm), titolo=tit, locale=loc, codloc=cod,
                              cap=cap, tot_n=int(tg.group(1)), tot=float(tg.group(2)), mese=mese))

# capienza per spettacoli a zero vendite: valore del locale letto dalle altre pagine
caps = [s["cap"] for s in shows if s["cap"]]
default_cap = max(set(caps), key=caps.count) if caps else None
for s in shows:
    s["cap_derived"] = s["cap"] is None
    if s["cap"] is None:
        s["cap"] = default_cap
shows.sort(key=lambda s: (s["date"], s["time"]))
sales.sort(key=lambda s: (s["date"], s["time"]))

# ---- stato precedente ----
def descr_tipo(p, prezzo):
    return f"Auto da {p} · {prezzo:.0f} €"

# tariffe: [codice, descrizione, tipologia, persone per auto, da definire]
tariffe = []
storico = []  # [data, auto, incasso, persone]
prev_ids = []
if PREV and os.path.exists(PREV):
    wbp = load_workbook(PREV)
    for r in wbp["Tariffe"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            pers = int(r[3]) if isinstance(r[3], (int, float)) and r[3] > 0 else None
            tariffe.append([r[0], r[1], r[2], pers, r[1] == "DA DEFINIRE"])
    ws_old = wbp["Storico"]
    cols = {str(c.value).strip(): i for i, c in enumerate(ws_old[1]) if c.value}
    c_auto = cols.get("Auto totali", cols.get("Biglietti totali", 1))
    c_inc = cols.get("Incasso totale", 2)
    c_pers = cols.get("Persone totali")
    for r in ws_old.iter_rows(min_row=2, values_only=True):
        if r[0]:
            pers = r[c_pers] if c_pers is not None and isinstance(r[c_pers], (int, float)) else None
            storico.append([r[0], r[c_auto], r[c_inc], pers])
    prev_ids = [r[0] for r in wbp["Spettacoli"].iter_rows(min_row=2, values_only=True) if r[0]]

# persone per auto di ogni codice: prima il foglio Tariffe, poi la regola sul prezzo
prezzo_codice = {}
for s in sales:
    prezzo_codice.setdefault(s["code"], s["prezzo"])
known = {t[0]: t for t in tariffe}
new_codes = []
for code, prezzo in prezzo_codice.items():
    p = PERSONE_PER_PREZZO.get(round(prezzo, 2))
    t = known.get(code)
    if t is None:
        t = [code, descr_tipo(p, prezzo) if p else "DA DEFINIRE", "Intero" if p else None, p, not p]
        tariffe.append(t); known[code] = t
        if not p:
            new_codes.append(code)
    else:
        if t[3] is None and p:
            t[3] = p
        if p and (not t[1] or t[1] == "DA DEFINIRE" or re.fullmatch(r"Tariffa \d+ €", str(t[1]))):
            t[1] = descr_tipo(p, prezzo)
        t[4] = t[3] is None
non_mappati = []
for s in sales:
    s["pers_auto"] = known[s["code"]][3]
    s["persone"] = s["n"] * s["pers_auto"] if s["pers_auto"] else 0
    if not s["pers_auto"] and s["n"]:
        non_mappati.append(s)
cur_ids = {s["id"] for s in shows}
missing = [i for i in prev_ids if i not in cur_ids]

tot_n = sum(s["n"] for s in sales); tot_l = round(sum(s["lordo"] for s in sales), 2)
tot_p = sum(s["persone"] for s in sales)
# storico vecchio senza persone: con la regola dei 12 € a persona le persone sono incasso / 12
for r in storico:
    if r[3] is None and isinstance(r[2], (int, float)):
        r[3] = round(r[2] / EURO_A_PERSONA)
storico = [r for r in storico if r[0] != SNAP]
storico.append([SNAP, tot_n, tot_l, tot_p])

# ---- stili ----
HDR = PatternFill("solid", fgColor="1F2430"); HF = Font(bold=True, color="FFFFFF", name="Arial", size=10)
ALT = PatternFill("solid", fgColor="F2F3F5"); YEL = PatternFill("solid", fgColor="FFF2A8")
BF = Font(name="Arial", size=10); thin = Side(style="thin", color="D9DCE1")
EUR = '#,##0.00 "€"'; DATE = "DD/MM/YYYY"; TIME = "HH:MM"; PCT = "0.0%"; INT = "#,##0"

def header(ws, cols, widths, row=1):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = HDR; cell.font = HF; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

def body(ws, nrows, ncols, fmts, start=2):
    for r in range(start, start + nrows):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c); cell.font = BF; cell.border = Border(bottom=thin)
            if (r - start) % 2 == 1: cell.fill = ALT
            if fmts.get(c): cell.number_format = fmts[c]
    if nrows:
        ws.auto_filter.ref = f"A{start-1}:{ws.cell(row=start+nrows-1, column=ncols).coordinate}"

wb = Workbook()
dash = wb.active; dash.title = "Dashboard"
wsT = wb.create_sheet("Tariffe"); wsY = wb.create_sheet("Tipologie"); wsS = wb.create_sheet("Spettacoli")
wsV = wb.create_sheet("Vendite"); wsP = wb.create_sheet("Per tariffa"); wsR = wb.create_sheet("Per serata")
wsH = wb.create_sheet("Storico")

# Tariffe (modificabile: la colonna D decide quante persone vale ogni auto)
header(wsT, ["Codice titolo", "Descrizione", "Tipologia (Intero/Ridotto/Omaggio)", "Persone per auto"], [14, 26, 30, 18])
for i, t in enumerate(tariffe, 2):
    for c, v in enumerate(t[:4], 1): wsT.cell(row=i, column=c, value=v)
body(wsT, len(tariffe), 4, {})
for i, t in enumerate(tariffe, 2):
    if t[4]:
        for c in range(1, 5): wsT.cell(row=i, column=c).fill = YEL
NT = len(tariffe) + 1
TR = "Tariffe!$A$2:$A$200"

# Vendite: una riga per spettacolo e tariffa. Ogni titolo è un'auto.
nV = len(sales)
header(wsV, ["ID spettacolo", "Data", "Ora", "Settore", "Codice titolo", "Descrizione tariffa", "Prezzo unitario",
             "N° auto (titoli emessi)", "N° annullati", "Incasso lordo", "Prevendita", "Imponibile IVA", "IVA",
             "Mese C1 di origine", "Persone per auto", "Persone"],
       [16, 12, 8, 12, 12, 22, 13, 12, 11, 13, 12, 14, 11, 12, 12, 10])
for i, s in enumerate(sales, 2):
    vals = [s["id"], s["date"], s["time"], s["settore"], s["code"],
            f'=IFERROR(INDEX(Tariffe!$B$2:$B$200,MATCH(E{i},{TR},0)),"DA DEFINIRE")',
            s["prezzo"], s["n"], s["ann"], s["lordo"], s["prev"], s["impiva"], s["iva"], s["mese"],
            f'=IFERROR(INDEX(Tariffe!$D$2:$D$200,MATCH(E{i},{TR},0))*1,0)', f"=H{i}*O{i}"]
    for c, v in enumerate(vals, 1): wsV.cell(row=i, column=c, value=v)
body(wsV, nV, 16, {2: DATE, 3: TIME, 7: EUR, 8: INT, 9: INT, 10: EUR, 11: EUR, 12: EUR, 13: EUR, 15: INT, 16: INT})
V = lambda col: f"Vendite!${col}$2:${col}$1000"

# Tipologie: auto da 2, 3, 4 persone
header(wsY, ["Tipologia", "Prezzo", "Persone per auto", "N° auto", "Persone", "Incasso", "% delle auto"], [22, 10, 16, 10, 10, 13, 13])
tipi_rows = [(f"Auto da {p} persone", p * EURO_A_PERSONA, p) for p in TIPI]
if non_mappati:
    tipi_rows.append(("Da definire (prezzo non previsto)", None, 0))
NY = len(tipi_rows) + 1
for i, (lab, prezzo, p) in enumerate(tipi_rows, 2):
    wsY.cell(row=i, column=1, value=lab); wsY.cell(row=i, column=2, value=prezzo); wsY.cell(row=i, column=3, value=p)
    wsY.cell(row=i, column=4, value=f"=SUMIF({V('O')},C{i},{V('H')})")
    wsY.cell(row=i, column=5, value=f"=SUMIF({V('O')},C{i},{V('P')})")
    wsY.cell(row=i, column=6, value=f"=SUMIF({V('O')},C{i},{V('J')})")
    wsY.cell(row=i, column=7, value=f"=IF(SUM($D$2:$D${NY})>0,D{i}/SUM($D$2:$D${NY}),0)")
body(wsY, NY - 1, 7, {2: EUR, 3: INT, 4: INT, 5: INT, 6: EUR, 7: PCT})
yr = NY + 1
wsY.cell(row=yr, column=1, value="Totale")
for c, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
    wsY.cell(row=yr, column=c, value=f"=SUM({col}2:{col}{NY})")
for c in range(1, 8):
    x = wsY.cell(row=yr, column=c); x.font = Font(name="Arial", size=10, bold=True); x.border = Border(top=Side(style="thin", color="1F2430"))
for c, f in ((4, INT), (5, INT), (6, EUR), (7, PCT)): wsY.cell(row=yr, column=c).number_format = f

# Spettacoli
nS = len(shows); SL = nS + 1
header(wsS, ["ID", "Data", "Giorno settimana", "Ora", "Capienza (auto)", "Auto", "Persone", "Auto da 2", "Auto da 3",
             "Auto da 4", "Annullati", "Incasso lordo", "Imponibile IVA", "IVA", "Occupazione % (auto)", "Stato", "Chiave ordinamento"],
       [16, 12, 14, 8, 11, 9, 10, 10, 10, 10, 10, 13, 14, 11, 13, 14, 10])
for i, s in enumerate(shows, 2):
    vals = [s["id"], s["date"], '='+'CHOOSE(WEEKDAY(B,2),"Lunedì","Martedì","Mercoledì","Giovedì","Venerdì","Sabato","Domenica")'.replace('B,','B'+str(i)+','), s["time"], s["cap"],
            f"=SUMIF({V('A')},A{i},{V('H')})", f"=SUMIF({V('A')},A{i},{V('P')})"] + \
           [f"=SUMIFS({V('H')},{V('A')},A{i},{V('O')},{p})" for p in TIPI] + \
           [f"=SUMIF({V('A')},A{i},{V('I')})", f"=SUMIF({V('A')},A{i},{V('J')})", f"=SUMIF({V('A')},A{i},{V('L')})",
            f"=SUMIF({V('A')},A{i},{V('M')})", f'=IF(E{i}>0,F{i}/E{i},0)',
            f'=IF(F{i}>0,"Venduto","Zero vendite")', f"=F{i}+L{i}/100000+(1000-ROW())/100000000"]
    for c, v in enumerate(vals, 1): wsS.cell(row=i, column=c, value=v)
body(wsS, nS, 17, {2: DATE, 4: TIME, 5: INT, 6: INT, 7: INT, 8: INT, 9: INT, 10: INT, 11: INT, 12: EUR, 13: EUR, 14: EUR, 15: PCT})
for i, s in enumerate(shows, 2):
    if s["cap_derived"]:
        wsS.cell(row=i, column=5).font = Font(name="Arial", size=10, italic=True, color="6B7280")
wsS.column_dimensions["Q"].hidden = True
wsS.conditional_formatting.add(f"P2:P{SL}", FormulaRule(formula=[f'P2="Venduto"'], font=Font(color="1E6B3A", bold=True)))

# Per tariffa
header(wsP, ["Codice titolo", "Descrizione", "Persone per auto", "N° auto", "Persone", "Incasso", "% sull'incasso totale"], [14, 26, 15, 10, 10, 14, 18])
for i in range(2, NT + 1):
    wsP.cell(row=i, column=1, value=f"=Tariffe!A{i}")
    wsP.cell(row=i, column=2, value=f"=Tariffe!B{i}")
    wsP.cell(row=i, column=3, value=f"=Tariffe!D{i}")
    wsP.cell(row=i, column=4, value=f"=SUMIF({V('E')},A{i},{V('H')})")
    wsP.cell(row=i, column=5, value=f"=SUMIF({V('E')},A{i},{V('P')})")
    wsP.cell(row=i, column=6, value=f"=SUMIF({V('E')},A{i},{V('J')})")
    wsP.cell(row=i, column=7, value=f"=IF(SUM($F$2:$F${NT})>0,F{i}/SUM($F$2:$F${NT}),0)")
tr = NT + 1
wsP.cell(row=tr, column=1, value="Totale")
for c, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
    wsP.cell(row=tr, column=c, value=f"=SUM({col}2:{col}{NT})")
body(wsP, NT - 1, 7, {3: INT, 4: INT, 5: INT, 6: EUR, 7: PCT})
for c in range(1, 8):
    x = wsP.cell(row=tr, column=c); x.font = Font(name="Arial", size=10, bold=True); x.border = Border(top=Side(style="thin", color="1F2430"))
for c, f in ((4, INT), (5, INT), (6, EUR), (7, PCT)): wsP.cell(row=tr, column=c).number_format = f

# Per serata: auto e persone per turno (fasce: 23:59 include l'orario 23:58 presente nel C1 del 23/10)
slots = [("20:15", "20:00", "21:00"), ("21:15", "21:00", "22:00"), ("22:15", "22:00", "23:00"), ("23:15", "23:00", "23:30"), ("23:59", "23:30", "24:00")]
dates = sorted({s["date"] for s in shows}); ND = len(dates)
NSL = len(slots)
A0, P0, Y0 = 6, 6 + NSL, 6 + 2 * NSL          # prime colonne di auto per turno, persone per turno, auto per tipologia
LBL = Y0 + len(TIPI)                           # colonna etichette del grafico
header(wsR, ["Data", "Giorno", "Auto totali", "Persone totali", "Incasso totale"] + [f"Auto {s[0]}" for s in slots]
       + [f"Persone {s[0]}" for s in slots] + [f"Auto da {p}" for p in TIPI],
       [12, 13, 11, 13, 14] + [10] * NSL + [11] * NSL + [10] * len(TIPI))
col = lambda c: wsR.cell(row=1, column=c).column_letter
for i, d in enumerate(dates, 2):
    wsR.cell(row=i, column=1, value=d)
    wsR.cell(row=i, column=2, value='='+'CHOOSE(WEEKDAY(A,2),"Lunedì","Martedì","Mercoledì","Giovedì","Venerdì","Sabato","Domenica")'.replace('A,','A'+str(i)+','))
    wsR.cell(row=i, column=3, value=f"=SUM({col(A0)}{i}:{col(A0+NSL-1)}{i})")
    wsR.cell(row=i, column=4, value=f"=SUM({col(P0)}{i}:{col(P0+NSL-1)}{i})")
    wsR.cell(row=i, column=5, value=f"=SUMIF({V('B')},A{i},{V('J')})")
    for k, (lab, lo, hi) in enumerate(slots):
        cond = f'{V("C")},">="&TIMEVALUE("{lo}")'
        if hi != "24:00":
            cond += f',{V("C")},"<"&TIMEVALUE("{hi}")'
        wsR.cell(row=i, column=A0 + k, value=f'=SUMIFS({V("H")},{V("B")},A{i},{cond})')
        wsR.cell(row=i, column=P0 + k, value=f'=SUMIFS({V("P")},{V("B")},A{i},{cond})')
    for k, p in enumerate(TIPI):
        wsR.cell(row=i, column=Y0 + k, value=f'=SUMIFS({V("H")},{V("B")},A{i},{V("O")},{p})')
    wsR.cell(row=i, column=LBL, value=f'=TEXT(A{i},"DD/MM")')
wsR.cell(row=1, column=LBL, value="Etichetta grafico")
body(wsR, ND, LBL - 1, {1: DATE, 3: INT, 4: INT, 5: EUR, **{c: INT for c in range(A0, LBL)}})
for c0 in (A0, P0):
    wsR.conditional_formatting.add(f"{col(c0)}2:{col(c0+NSL-1)}{ND+1}", ColorScaleRule(start_type="min", start_color="FFFFFF", mid_type="percentile", mid_value=50, mid_color="F4B6A6", end_type="max", end_color="B3261E"))

# Storico (valori di snapshot, mai formule: devono restare congelati)
header(wsH, ["Data estrazione", "Auto totali", "Persone totali", "Incasso totale", "Delta auto", "Delta persone", "Delta incasso"], [16, 12, 14, 15, 12, 14, 15])
for i, r in enumerate(storico, 2):
    wsH.cell(row=i, column=1, value=r[0]); wsH.cell(row=i, column=2, value=r[1])
    wsH.cell(row=i, column=3, value=r[3]); wsH.cell(row=i, column=4, value=r[2])
    for c, a in ((5, "B"), (6, "C"), (7, "D")):
        wsH.cell(row=i, column=c, value="n.d." if i == 2 else f"={a}{i}-{a}{i-1}")
DINT = '+#,##0;-#,##0;0'
body(wsH, len(storico), 7, {1: "DD/MM/YYYY HH:MM", 2: INT, 3: INT, 4: EUR, 5: DINT, 6: DINT, 7: '+#,##0.00 "€";-#,##0.00 "€";0,00 "€"'})
for i in range(2, len(storico) + 2):
    for c in (5, 6, 7): wsH.cell(row=i, column=c).alignment = Alignment(horizontal="right")
HL = len(storico) + 1
for i in range(2, HL + 1):
    wsH.cell(row=i, column=8, value=f'=TEXT(A{i},"DD/MM HH:MM")')
wsH.cell(row=1, column=8, value="Etichetta grafico")

# Dashboard
dash.sheet_view.showGridLines = False
for c, w in zip("ABCDEFGHIJKLMNO", [2, 22, 18, 3, 22, 18, 3, 22, 18, 3, 12, 8, 9, 10, 13]):
    dash.column_dimensions[c].width = w
dash["B2"] = "HORROR DAYS · Vendite biglietti"; dash["B2"].font = Font(name="Arial", size=18, bold=True, color="1F2430")
dash["B3"] = "Drive In Pozzuoli · MIX MARK S.R.L. · fonte: Riepiloghi mensili C1 (18Tickets) · 1 biglietto = 1 auto"; dash["B3"].font = Font(name="Arial", size=10, color="6B7280")
tipi_txt = '&" · "&'.join(f'"Da {p}: "&Tipologie!D{k}' for k, p in enumerate(TIPI, 2))
kpis = [
    ("Incasso lordo totale", f"=SUM({V('J')})", EUR),
    ("Auto vendute", f"=SUM({V('H')})", INT),
    ("Persone", f"=SUM({V('P')})", INT),
    ("Persone per auto (media)", f"=IF(SUM({V('H')})>0,SUM({V('P')})/SUM({V('H')}),0)", "0.00"),
    ("Occupazione media (auto)", f"=IF(SUM(Spettacoli!E2:E{SL})>0,SUM(Spettacoli!F2:F{SL})/SUM(Spettacoli!E2:E{SL}),0)", "0.00%"),
    ("Spettacoli con vendite / totali", f'=COUNTIF(Spettacoli!P2:P{SL},"Venduto")&" / "&COUNTA(Spettacoli!A2:A{SL})', "@"),
    ("Auto per tipologia", "=" + tipi_txt, "@"),
    ("Prezzo medio per auto", f"=IF(SUM({V('H')})>0,SUM({V('J')})/SUM({V('H')}),0)", EUR),
    ("Ultimo aggiornamento", f"=MAX(Storico!A2:A{HL})", "DD/MM/YYYY HH:MM"),
]
pos = [(c, r) for r in (5, 8, 11) for c in ("B", "E", "H")]
card = PatternFill("solid", fgColor="F2F3F5")
for (lab, f, fmt), (c, r) in zip(kpis, pos):
    c2 = chr(ord(c) + 1)
    dash.merge_cells(f"{c}{r}:{c2}{r}"); dash.merge_cells(f"{c}{r+1}:{c2}{r+1}")
    a = dash[f"{c}{r}"]; a.value = lab; a.font = Font(name="Arial", size=9, color="6B7280")
    b = dash[f"{c}{r+1}"]; b.value = f; b.number_format = fmt; b.font = Font(name="Arial", size=16 if fmt != "@" else 13, bold=True, color="1F2430")
    b.alignment = Alignment(horizontal="left")
    for rr in (r, r + 1):
        for cc in (c, c2): dash[f"{cc}{rr}"].fill = card
    dash.row_dimensions[r + 1].height = 26

# Top 5
dash["K4"] = "Top 5 spettacoli per auto vendute"; dash["K4"].font = Font(name="Arial", size=11, bold=True, color="1F2430")
for j, h in enumerate(["Data", "Ora", "Auto", "Persone", "Incasso"]):
    c = dash.cell(row=5, column=11 + j, value=h); c.fill = HDR; c.font = HF; c.alignment = Alignment(horizontal="center")
for k in range(1, 6):
    r = 5 + k
    m = f"MATCH(LARGE(Spettacoli!$Q$2:$Q${SL},{k}),Spettacoli!$Q$2:$Q${SL},0)"
    dash.cell(row=r, column=11, value=f"=INDEX(Spettacoli!$B$2:$B${SL},{m})").number_format = DATE
    dash.cell(row=r, column=12, value=f"=INDEX(Spettacoli!$D$2:$D${SL},{m})").number_format = TIME
    dash.cell(row=r, column=13, value=f"=INDEX(Spettacoli!$F$2:$F${SL},{m})").number_format = INT
    dash.cell(row=r, column=14, value=f"=INDEX(Spettacoli!$G$2:$G${SL},{m})").number_format = INT
    dash.cell(row=r, column=15, value=f"=INDEX(Spettacoli!$L$2:$L${SL},{m})").number_format = EUR
    for c in range(11, 16):
        x = dash.cell(row=r, column=c); x.font = BF; x.border = Border(bottom=thin)
        if k % 2 == 0: x.fill = ALT

# Grafici
ch1 = BarChart(); ch1.type = "col"; ch1.grouping = "stacked"; ch1.overlap = 100
ch1.title = "Auto per serata e tipologia"; ch1.style = 2
ch1.add_data(Reference(wsR, min_col=Y0, max_col=Y0 + len(TIPI) - 1, min_row=1, max_row=ND + 1), titles_from_data=True)
ch1.set_categories(Reference(wsR, min_col=LBL, min_row=2, max_row=ND + 1))
for sr, colr in zip(ch1.series, ["F4B6A6", "D43A31", "8B1E1E"]):
    sr.graphicalProperties.solidFill = colr
ch1.y_axis.numFmt = "0"; ch1.height = 7.5; ch1.width = 16
ch1.y_axis.delete = False; ch1.x_axis.delete = False
ch1.plotVisOnly = False
dash.add_chart(ch1, "B15")

ch2 = PieChart(); ch2.title = "Auto per tipologia"
ch2.add_data(Reference(wsY, min_col=4, min_row=1, max_row=NY), titles_from_data=True)
ch2.set_categories(Reference(wsY, min_col=1, min_row=2, max_row=NY))
ch2.dataLabels = DataLabelList(); ch2.dataLabels.showPercent = True; ch2.height = 7.5; ch2.width = 11
ch2.plotVisOnly = False
dash.add_chart(ch2, "H15")

ch3 = LineChart(); ch3.title = "Andamento vendite (persone cumulate)"; ch3.style = 2
ch3.add_data(Reference(wsH, min_col=3, min_row=1, max_row=HL), titles_from_data=True)
ch3.set_categories(Reference(wsH, min_col=8, min_row=2, max_row=HL))
ch3.legend = None; ch3.y_axis.numFmt = "0"; ch3.x_axis.number_format = "DD/MM"
ch3.series[0].graphicalProperties.line.solidFill = "1F2430"; ch3.series[0].marker.symbol = "circle"
ch3.height = 7.5; ch3.width = 16; ch3.y_axis.delete = False; ch3.x_axis.delete = False
ch3.plotVisOnly = False
dash.add_chart(ch3, "B31")

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = "1F2430" if ws.title == "Dashboard" else None
os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
wb.save(OUT)

# ---- controlli ----
checks = []
def chk(name, ok, detail=""):
    checks.append({"controllo": name, "ok": bool(ok), "dettaglio": detail})
chk("Incasso Vendite = somma TOTALE GENERALE", abs(tot_l - pdf_tot_gen) < 0.01, f"{tot_l:.2f} vs {pdf_tot_gen:.2f}")
chk("Spettacoli = pagine dei PDF", len(shows) == total_pages, f"{len(shows)} spettacoli, {total_pages} pagine")
bad = [s for s in sales if abs(s["impiva"] + s["iva"] - s["lordo"]) > 0.01]
chk("Imponibile IVA + IVA = incasso lordo per riga", not bad, f"{len(bad)} righe non quadrate")
annullati = [(s["id"], s["ann"]) for s in sales if s["ann"]]
prevendita = [(s["id"], s["prev"]) for s in sales if s["prev"]]
anomalies = []
if annullati: anomalies.append("Titoli annullati: " + ", ".join(f"{a} ({n})" for a, n in annullati))
if prevendita: anomalies.append("Prevendita diversa da zero: " + ", ".join(f"{a} ({v:.2f} €)" for a, v in prevendita))
if new_codes: anomalies.append("Codici tariffa nuovi da definire: " + ", ".join(new_codes))
if missing: anomalies.append("Spettacoli presenti prima e assenti nei nuovi C1: " + ", ".join(map(str, missing)))
odd = sorted({s["time"].strftime("%H:%M") for s in shows} - {"20:15", "21:15", "22:15", "23:15", "23:59", "00:15"})
if odd: anomalies.append("Orari non standard nei C1: " + ", ".join(odd))
derived = sum(1 for s in shows if s["cap_derived"])
if derived: anomalies.append(f"Capienza ricavata dalle altre pagine per {derived} spettacoli senza vendite ({default_cap})")
if non_mappati:
    anomalies.append("Auto senza numero di persone (prezzo non previsto, da indicare nel foglio Tariffe): " +
                     ", ".join(f"{s['id']} {s['code']} {s['prezzo']:.2f} € ({s['n']})" for s in non_mappati))
elif abs(tot_p * EURO_A_PERSONA - tot_l) > 0.01:
    anomalies.append(f"Persone × {EURO_A_PERSONA:.0f} € ({tot_p * EURO_A_PERSONA:.2f} €) diverso dall'incasso ({tot_l:.2f} €): controlla le persone per auto nel foglio Tariffe")

tmap = {t[0]: t for t in tariffe}
def slot(t):
    return "23:59" if t >= dt.time(23, 30) else t.strftime("%H:%M")
def per_tipo(rows):
    return {str(p): sum(v["n"] for v in rows if v["pers_auto"] == p) for p in TIPI}
def show_data(s):
    rows = [v for v in sales if v["id"] == s["id"]]
    auto = sum(v["n"] for v in rows)
    return {"id": s["id"], "data": s["date"].isoformat(), "ora": s["time"].strftime("%H:%M"), "slot": slot(s["time"]),
            "capienza": s["cap"], "auto": auto, "biglietti": auto, "persone": sum(v["persone"] for v in rows),
            "tipi": per_tipo(rows), "incasso": round(sum(v["lordo"] for v in rows), 2),
            "annullati": sum(v["ann"] for v in rows)}
data = {
    "aggiornato": SNAP.strftime("%Y-%m-%dT%H:%M"),
    "fonte": [os.path.basename(p) for p in pdfs],
    # 1 biglietto = 1 auto; la capienza è in auto
    "totali": {"incasso": tot_l, "auto": tot_n, "biglietti": tot_n, "persone": tot_p, "tipi": per_tipo(sales),
               "spettacoli": len(shows), "spettacoli_venduti": len({s["id"] for s in sales if s["n"]}),
               "capienza": sum(s["cap"] or 0 for s in shows)},
    "tipologie": [{"persone": p, "prezzo": p * EURO_A_PERSONA, "etichetta": f"Auto da {p}",
                   "auto": sum(v["n"] for v in sales if v["pers_auto"] == p),
                   "persone_tot": sum(v["persone"] for v in sales if v["pers_auto"] == p),
                   "incasso": round(sum(v["lordo"] for v in sales if v["pers_auto"] == p), 2)} for p in TIPI],
    "spettacoli": [show_data(s) for s in shows],
    "tariffe": [{"codice": c, "descrizione": t[1] or c, "persone_per_auto": t[3],
                 "auto": sum(v["n"] for v in sales if v["code"] == c),
                 "biglietti": sum(v["n"] for v in sales if v["code"] == c),
                 "persone": sum(v["persone"] for v in sales if v["code"] == c),
                 "incasso": round(sum(v["lordo"] for v in sales if v["code"] == c), 2)} for c, t in tmap.items()],
    "storico": [{"t": (r[0].strftime("%Y-%m-%dT%H:%M") if isinstance(r[0], dt.datetime) else str(r[0])),
                 "auto": r[1], "biglietti": r[1], "persone": r[3], "incasso": r[2]} for r in storico],
    "controlli": checks,
    "anomalie": anomalies,
}
os.makedirs(os.path.dirname(args.json) or ".", exist_ok=True)
with open(args.json, "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=1)
if args.xlsx_copy:
    import shutil; shutil.copyfile(OUT, args.xlsx_copy)

print(f"Incasso {tot_l:.2f} € · auto {tot_n} · persone {tot_p} · spettacoli {len(shows)} ({data['totali']['spettacoli_venduti']} con vendite)")
for c in checks: print(("OK  " if c["ok"] else "KO  ") + c["controllo"] + " · " + c["dettaglio"])
for a in anomalies: print("!!  " + a)
if not all(c["ok"] for c in checks):
    sys.exit(2)
